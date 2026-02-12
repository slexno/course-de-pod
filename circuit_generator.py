#!/usr/bin/env python3
"""Générateur de circuit type Formule 1.

- Demande le nombre de virages rapides/lents, chicanes et épingles.
- Génère un tracé fermé qui revient sur la ligne droite des stands.
- Évite les tracés auto-chevauchants.
- Exporte une image PNG du circuit.
- Met à jour le fichier Excel `circuit.xlsx` avec le détail de chaque section.
"""

from __future__ import annotations

import math
import random
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Tuple

from openpyxl import Workbook, load_workbook


@dataclass
class Segment:
    index: int
    type_segment: str
    categorie: str
    longueur_m: float
    angle_deg: float
    debut_x: float
    debut_y: float
    fin_x: float
    fin_y: float


TURN_LIBRARY = {
    "virage_rapide": {"angle": (25, 45), "radius": (90, 150)},
    "virage_lent": {"angle": (55, 95), "radius": (35, 75)},
    "chicane": {"angle": (28, 45), "radius": (30, 55)},
    "epingle": {"angle": (130, 170), "radius": (18, 28)},
}


def ask_int(prompt: str) -> int:
    while True:
        raw = input(prompt).strip()
        try:
            value = int(raw)
            if value < 0:
                raise ValueError
            return value
        except ValueError:
            print("Veuillez entrer un entier positif (0 accepté).")


def generate_turn_segments(counts: dict[str, int]) -> List[str]:
    turns: List[str] = []
    for key, count in counts.items():
        turns.extend([key] * count)
    random.shuffle(turns)
    return turns


def points_equal(a: Tuple[float, float], b: Tuple[float, float], eps: float = 1e-6) -> bool:
    return abs(a[0] - b[0]) <= eps and abs(a[1] - b[1]) <= eps


def orient(a: Tuple[float, float], b: Tuple[float, float], c: Tuple[float, float]) -> float:
    return (b[0] - a[0]) * (c[1] - a[1]) - (b[1] - a[1]) * (c[0] - a[0])


def on_segment(a: Tuple[float, float], b: Tuple[float, float], c: Tuple[float, float], eps: float = 1e-9) -> bool:
    return (
        min(a[0], b[0]) - eps <= c[0] <= max(a[0], b[0]) + eps
        and min(a[1], b[1]) - eps <= c[1] <= max(a[1], b[1]) + eps
        and abs(orient(a, b, c)) <= eps
    )


def segments_intersect(
    p1: Tuple[float, float],
    q1: Tuple[float, float],
    p2: Tuple[float, float],
    q2: Tuple[float, float],
) -> bool:
    o1 = orient(p1, q1, p2)
    o2 = orient(p1, q1, q2)
    o3 = orient(p2, q2, p1)
    o4 = orient(p2, q2, q1)

    if (o1 > 0 > o2 or o1 < 0 < o2) and (o3 > 0 > o4 or o3 < 0 < o4):
        return True

    if on_segment(p1, q1, p2) or on_segment(p1, q1, q2) or on_segment(p2, q2, p1) or on_segment(p2, q2, q1):
        return True

    return False


def has_self_intersection(points: List[Tuple[float, float]]) -> bool:
    if len(points) < 4:
        return False

    lines = list(zip(points[:-1], points[1:]))
    n = len(lines)

    for i in range(n):
        a1, a2 = lines[i]
        for j in range(i + 1, n):
            b1, b2 = lines[j]

            if abs(i - j) <= 1:
                continue
            if i == 0 and j == n - 1:
                continue

            shared_endpoint = (
                points_equal(a1, b1) or points_equal(a1, b2) or points_equal(a2, b1) or points_equal(a2, b2)
            )
            if shared_endpoint:
                continue

            if segments_intersect(a1, a2, b1, b2):
                return True

    return False


def add_straight(
    segments: List[Segment],
    points: List[Tuple[float, float]],
    heading: float,
    length: float,
    category: str = "ligne_droite",
) -> None:
    x0, y0 = points[-1]
    x1 = x0 + length * math.cos(math.radians(heading))
    y1 = y0 + length * math.sin(math.radians(heading))

    segments.append(
        Segment(
            index=len(segments) + 1,
            type_segment="ligne_droite",
            categorie=category,
            longueur_m=round(length, 2),
            angle_deg=0.0,
            debut_x=round(x0, 2),
            debut_y=round(y0, 2),
            fin_x=round(x1, 2),
            fin_y=round(y1, 2),
        )
    )
    points.append((x1, y1))


def add_arc(
    segments: List[Segment],
    points: List[Tuple[float, float]],
    heading: float,
    turn_type: str,
    direction: int,
) -> float:
    profile = TURN_LIBRARY[turn_type]
    angle = random.uniform(*profile["angle"])
    radius = random.uniform(*profile["radius"])

    sweep = direction * angle
    arc_len = radius * math.radians(angle)

    x0, y0 = points[-1]

    cx = x0 - direction * radius * math.sin(math.radians(heading))
    cy = y0 + direction * radius * math.cos(math.radians(heading))

    start_theta = math.degrees(math.atan2(y0 - cy, x0 - cx))

    steps = max(8, int(abs(angle) // 3))
    for i in range(1, steps + 1):
        t = start_theta + sweep * (i / steps)
        x = cx + radius * math.cos(math.radians(t))
        y = cy + radius * math.sin(math.radians(t))
        points.append((x, y))

    x1, y1 = points[-1]

    segments.append(
        Segment(
            index=len(segments) + 1,
            type_segment="virage",
            categorie=turn_type,
            longueur_m=round(arc_len, 2),
            angle_deg=round(sweep, 2),
            debut_x=round(x0, 2),
            debut_y=round(y0, 2),
            fin_x=round(x1, 2),
            fin_y=round(y1, 2),
        )
    )

    return heading + sweep


def build_track_once(counts: dict[str, int]) -> Tuple[List[Segment], List[Tuple[float, float]]]:
    turns = generate_turn_segments(counts)
    segments: List[Segment] = []
    points: List[Tuple[float, float]] = [(0.0, 0.0)]

    heading = 0.0
    pit_length = random.uniform(220, 360)
    add_straight(segments, points, heading, pit_length, category="ligne_des_stands")

    for i, turn_type in enumerate(turns):
        if turn_type == "chicane":
            dir_a = random.choice([-1, 1])
            heading = add_arc(segments, points, heading, turn_type, dir_a)
            short_straight = random.uniform(45, 90)
            add_straight(segments, points, heading, short_straight, category="liaison_chicane")
            heading = add_arc(segments, points, heading, turn_type, -dir_a)
        else:
            direction = random.choice([-1, 1])
            heading = add_arc(segments, points, heading, turn_type, direction)

        if i < len(turns) - 1:
            straight_len = random.uniform(80, 260)
            add_straight(segments, points, heading, straight_len)

    x_end, y_end = points[-1]
    x_start, y_start = points[0]
    dx, dy = x_start - x_end, y_start - y_end
    closure_len = math.hypot(dx, dy)

    if closure_len > 1:
        heading_to_start = math.degrees(math.atan2(dy, dx))
        add_straight(segments, points, heading_to_start, closure_len, category="retour_stands")

    return segments, points


def build_track(counts: dict[str, int], max_attempts: int = 250) -> Tuple[List[Segment], List[Tuple[float, float]]]:
    for _ in range(max_attempts):
        segments, points = build_track_once(counts)
        if not has_self_intersection(points):
            return segments, points

    return build_track_once(counts)


def export_image(points: List[Tuple[float, float]], output: Path) -> None:
    import matplotlib.pyplot as plt

    xs, ys = zip(*points)
    plt.figure(figsize=(10, 8))
    plt.plot(xs, ys, linewidth=7, color="#222")
    plt.plot(xs, ys, linewidth=4, color="#bdbdbd")
    plt.scatter([xs[0]], [ys[0]], s=120, marker="s", color="red", label="Ligne des stands")
    plt.axis("equal")
    plt.axis("off")
    plt.title("Circuit généré automatiquement")
    plt.legend(loc="upper right")
    plt.tight_layout()
    plt.savefig(output, dpi=180)
    plt.close()


def update_excel(segments: List[Segment], excel_path: Path) -> Path:
    if excel_path.exists():
        try:
            wb = load_workbook(excel_path)
        except PermissionError:
            wb = Workbook()
    else:
        wb = Workbook()

    if "Circuit" in wb.sheetnames:
        ws = wb["Circuit"]
        ws.delete_rows(1, ws.max_row)
    else:
        ws = wb.create_sheet("Circuit")

    headers = [
        "Index",
        "Type",
        "Categorie",
        "Longueur_m",
        "Angle_deg",
        "Debut_X",
        "Debut_Y",
        "Fin_X",
        "Fin_Y",
    ]
    ws.append(headers)

    for seg in segments:
        ws.append(
            [
                seg.index,
                seg.type_segment,
                seg.categorie,
                seg.longueur_m,
                seg.angle_deg,
                seg.debut_x,
                seg.debut_y,
                seg.fin_x,
                seg.fin_y,
            ]
        )

    try:
        wb.save(excel_path)
        return excel_path
    except PermissionError:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        fallback_path = excel_path.with_name(f"{excel_path.stem}_{timestamp}.xlsx")
        wb.save(fallback_path)
        print(
            "[INFO] Impossible d'écrire dans 'circuit.xlsx' (fichier probablement ouvert). "
            f"Export réalisé dans: {fallback_path}"
        )
        return fallback_path


def main() -> None:
    print("=== Générateur de circuit style Formule 1 ===")
    random.seed()

    counts = {
        "virage_rapide": ask_int("Nombre de virages rapides : "),
        "virage_lent": ask_int("Nombre de virages lents : "),
        "chicane": ask_int("Nombre de chicanes : "),
        "epingle": ask_int("Nombre d'épingles : "),
    }

    if sum(counts.values()) == 0:
        print("Vous devez demander au moins un virage pour générer un circuit.")
        return

    segments, points = build_track(counts)

    image_path = Path("circuit.png")
    excel_path = Path("circuit.xlsx")

    export_image(points, image_path)
    excel_written_path = update_excel(segments, excel_path)

    print(f"Image créée : {image_path.resolve()}")
    print(f"Tableau Excel mis à jour : {excel_written_path.resolve()}")
    print(f"Nombre total de segments : {len(segments)}")


if __name__ == "__main__":
    main()