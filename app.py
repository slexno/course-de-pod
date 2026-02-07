from __future__ import annotations

import math
import random
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
CIRCUIT_FILE = BASE_DIR / "circuit.xlsx"
CIRCUIT_PNG = BASE_DIR / "circuit.png"

app = Flask(__name__, template_folder=str(BASE_DIR / "templates"), static_folder=str(BASE_DIR / "static"))


@dataclass
class Player:
    player_id: int
    name: str
    engine_power: int
    downforce: int
    dexterity: int
    aggressiveness: int
    tire_wear: float = 0.0
    status: str = "running"
    risk_penalty_sections: int = 0
    next_section_bonus: int = 0
    prep_defense_bonus: int = 0
    corner_exit_bonus: int = 0
    successful_overtakes: int = 0
    successful_defenses: int = 0
    malus_count: int = 0
    color: str = "#22c55e"


class GameState:
    def __init__(self) -> None:
        self.players: list[Player] = []
        self.positions: list[int] = []
        self.segment_index = 0
        self.active_position_index = 0
        self.last_duel: dict[str, Any] | None = None
        self.weather = "dry"
        self.current_lap = 1
        self.total_laps = 3

    @property
    def started(self) -> bool:
        return len(self.players) >= 3


game_state = GameState()


PLAYER_COLORS = [
    "#ef4444", "#3b82f6", "#22c55e", "#f59e0b", "#a855f7",
    "#06b6d4", "#f97316", "#84cc16", "#e879f9", "#14b8a6",
]


def color_for_player(idx: int) -> str:
    return PLAYER_COLORS[(idx - 1) % len(PLAYER_COLORS)]


def safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def dnd_modifier(stat: int) -> int:
    return max(-5, min(5, math.floor((stat - 10) / 2)))


def weather_dex_penalty() -> int:
    return {"dry": 0, "windy": -1, "light_rain": -2, "storm": -4}.get(game_state.weather, 0)


def load_circuit_segments() -> list[dict[str, Any]]:
    """Charge les segments Circuit en restant compatible avec variantes de colonnes."""
    if not CIRCUIT_FILE.exists():
        return []

    wb = load_workbook(CIRCUIT_FILE, data_only=True)
    if "Circuit" not in wb.sheetnames:
        return []

    ws = wb["Circuit"]
    segments: list[dict[str, Any]] = []

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    header_map: dict[str, int] = {}
    for i, col in enumerate(header_row):
        if col is None:
            continue
        header_map[str(col).strip().lower()] = i

    def pick(row: tuple[Any, ...], *names: str, default: Any = None) -> Any:
        for name in names:
            idx = header_map.get(name.lower())
            if idx is not None and idx < len(row):
                return row[idx]
        return default

    for fallback_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=1):
        idx_raw = pick(row, "index", "idx", default=fallback_idx)
        if idx_raw is None:
            continue

        start_x = pick(row, "debut_x", "start_x", "x0", default=0)
        start_y = pick(row, "debut_y", "start_y", "y0", default=0)
        end_x = pick(row, "fin_x", "end_x", "x1", default=start_x)
        end_y = pick(row, "fin_y", "end_y", "y1", default=start_y)

        segments.append(
            {
                "index": safe_int(idx_raw, fallback_idx),
                "type": pick(row, "type", "type_segment", default="inconnu") or "inconnu",
                "category": pick(row, "categorie", "category", "segment_category", default="inconnu") or "inconnu",
                "length_m": float(pick(row, "longueur_m", "length_m", default=0) or 0),
                "start_x": float(start_x or 0),
                "start_y": float(start_y or 0),
                "end_x": float(end_x or 0),
                "end_y": float(end_y or 0),
            }
        )

    # fallback géométrie si aucune coordonnée exploitable
    if segments and all(
        s["start_x"] == 0 and s["start_y"] == 0 and s["end_x"] == 0 and s["end_y"] == 0
        for s in segments
    ):
        x = 0.0
        for s in segments:
            length = max(30.0, s["length_m"] or 30.0)
            s["start_x"], s["start_y"] = x, 0.0
            x += length
            s["end_x"], s["end_y"] = x, 0.0

    return segments


def wear_penalty(player: Player) -> int:
    if game_state.total_laps <= 0:
        return 0
    max_wear = game_state.total_laps * 30.0
    ratio = player.tire_wear / max_wear
    if ratio >= 1.0:
        return -4
    if ratio >= 0.75:
        return -3
    if ratio >= 0.5:
        return -2
    if ratio >= 0.3:
        return -1
    return 0


def current_segment(segments: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not segments:
        return None
    return segments[game_state.segment_index]


def get_player_by_id(pid: int) -> Player | None:
    return next((p for p in game_state.players if p.player_id == pid), None)


def get_active_player() -> Player | None:
    if not game_state.positions:
        return None
    if game_state.active_position_index >= len(game_state.positions):
        game_state.active_position_index = 0
    return get_player_by_id(game_state.positions[game_state.active_position_index])


def get_target_player() -> Player | None:
    if game_state.active_position_index <= 0:
        return None
    return get_player_by_id(game_state.positions[game_state.active_position_index - 1])


def attack_modifier(attacker: Player, segment: dict[str, Any], dangerous: bool) -> tuple[int, list[str]]:
    notes: list[str] = []
    w_pen = weather_dex_penalty()

    if segment["type"] == "ligne_droite":
        seg_mod = dnd_modifier(attacker.engine_power)
        notes.append(f"Ligne droite: engine {seg_mod:+d}")
    else:
        seg_mod = dnd_modifier(attacker.downforce) + dnd_modifier(attacker.aggressiveness)
        notes.append(
            f"Virage: downforce {dnd_modifier(attacker.downforce):+d}, agressivité {dnd_modifier(attacker.aggressiveness):+d}"
        )

    dex_mod = dnd_modifier(attacker.dexterity) + w_pen
    notes.append(f"Dex + météo: {dex_mod:+d}")

    total = seg_mod + dex_mod + attacker.next_section_bonus + wear_penalty(attacker)
    if segment["type"] == "ligne_droite":
        total += attacker.corner_exit_bonus
        if attacker.corner_exit_bonus:
            notes.append(f"Bonus sortie virage: {attacker.corner_exit_bonus:+d}")

    if dangerous:
        total += 2
        notes.append("Dépassement dangereux: +2")

    if attacker.next_section_bonus:
        notes.append(f"Bonus action prochaine section: {attacker.next_section_bonus:+d}")

    wp = wear_penalty(attacker)
    if wp:
        notes.append(f"Usure pneus: {wp:+d}")

    if attacker.risk_penalty_sections > 0:
        total -= 3
        notes.append(f"Malus risque: -3 ({attacker.risk_penalty_sections} sections restantes)")

    return total, notes


def defense_modifier(defender: Player) -> tuple[int, list[str]]:
    dex_mod = dnd_modifier(defender.dexterity) + weather_dex_penalty()
    total = dex_mod + defender.prep_defense_bonus + wear_penalty(defender)
    notes = [f"Défense dex+météo: {dex_mod:+d}"]

    if defender.prep_defense_bonus:
        notes.append(f"Préparation: {defender.prep_defense_bonus:+d}")

    wp = wear_penalty(defender)
    if wp:
        notes.append(f"Usure pneus: {wp:+d}")

    if defender.risk_penalty_sections > 0:
        total -= 3
        notes.append(f"Malus risque: -3 ({defender.risk_penalty_sections} sections restantes)")

    return total, notes


def dangerous_risk_penalty(attacker: Player, dangerous: bool) -> tuple[bool, str]:
    if not dangerous:
        return False, ""
    threshold = attacker.aggressiveness + dnd_modifier(attacker.dexterity)
    roll = random.randint(1, 20)
    if roll < threshold:
        attacker.risk_penalty_sections = max(attacker.risk_penalty_sections, 3)
        attacker.malus_count += 1
        return True, f"Risque pris: malus -3 pendant 3 sections ({roll} < {threshold})."
    return False, f"Risque évité ({roll} >= {threshold})."


def consume_turn(player: Player, segment: dict[str, Any], action_wear: float) -> None:
    player.tire_wear += action_wear
    if player.risk_penalty_sections > 0:
        player.risk_penalty_sections -= 1

    if segment["type"] != "ligne_droite":
        score = dnd_modifier(player.dexterity) + dnd_modifier(player.downforce)
        player.corner_exit_bonus = max(-2, min(3, score // 2))
    player.next_section_bonus = 0


def advance_turn() -> None:
    game_state.active_position_index += 1
    if game_state.active_position_index >= len(game_state.positions):
        game_state.active_position_index = 0
        segments = load_circuit_segments()
        segment_count = max(1, len(segments))
        game_state.segment_index = (game_state.segment_index + 1) % segment_count
        if game_state.segment_index == 0:
            game_state.current_lap += 1


def track_data(segments: list[dict[str, Any]]) -> dict[str, Any]:
    if not segments:
        return {"markers": [], "path": []}

    points: list[tuple[float, float]] = [(segments[0]["start_x"], segments[0]["start_y"]) ]
    points.extend((s["end_x"], s["end_y"]) for s in segments)

    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    min_x, max_x = min(xs), max(xs)
    min_y, max_y = min(ys), max(ys)
    span_x = max(max_x - min_x, 1e-6)
    span_y = max(max_y - min_y, 1e-6)

    def normalize(x: float, y: float) -> tuple[float, float]:
        x_pct = 4 + (92 * ((x - min_x) / span_x))
        y_pct = 96 - (92 * ((y - min_y) / span_y))
        return x_pct, y_pct

    path = [{"x": normalize(x, y)[0], "y": normalize(x, y)[1]} for x, y in points]

    seg = segments[game_state.segment_index]
    sx, sy = seg["start_x"], seg["start_y"]
    ex, ey = seg["end_x"], seg["end_y"]
    dx, dy = ex - sx, ey - sy
    norm = math.hypot(dx, dy) or 1.0
    # vecteur perpendiculaire pour espacer les pions latéralement sur la section
    px, py = -dy / norm, dx / norm

    markers = []
    count = max(1, len(game_state.positions))
    for i, pid in enumerate(game_state.positions):
        player = get_player_by_id(pid)
        if player is None:
            continue

        # rang => position le long de la section (plus proche sortie pour leader)
        t = 0.82 - (i * (0.6 / max(1, count - 1))) if count > 1 else 0.6
        # petit décalage latéral pour lisibilité
        lateral = (i - (count - 1) / 2) * 3.0
        x = sx + dx * t + px * lateral
        y = sy + dy * t + py * lateral

        mx, my = normalize(x, y)
        markers.append({
            "name": player.name,
            "x": mx,
            "y": my,
            "color": player.color,
            "status": "running",
        })

    return {"markers": markers, "path": path}


def qualification_order(segments: list[dict[str, Any]]) -> list[int]:
    scores: list[tuple[int, int]] = []
    for p in game_state.players:
        pace = 0
        corner_bonus = 0
        for seg in segments:
            if seg["type"] == "ligne_droite":
                pace += dnd_modifier(p.engine_power) + corner_bonus
            else:
                corner_bonus = max(-2, min(3, (dnd_modifier(p.dexterity) + dnd_modifier(p.downforce)) // 2))
                pace += dnd_modifier(p.downforce) + dnd_modifier(p.dexterity)
            pace += random.randint(1, 4)
        scores.append((pace, p.player_id))
    scores.sort(reverse=True)
    return [pid for _, pid in scores]


def final_stats() -> dict[str, Any]:
    ranked = [get_player_by_id(pid) for pid in game_state.positions]
    ranked = [p for p in ranked if p is not None]
    data = [
        {
            "position": i + 1,
            "name": p.name,
            "successful_overtakes": p.successful_overtakes,
            "successful_defenses": p.successful_defenses,
            "malus_count": p.malus_count,
        }
        for i, p in enumerate(ranked)
    ]
    return {"podium": data[:3], "others": data[3:], "full": data}


def serialize_state() -> dict[str, Any]:
    segments = load_circuit_segments()
    by_id = {p.player_id: p for p in game_state.players}
    ranking = [asdict(by_id[pid]) for pid in game_state.positions if pid in by_id]
    finished = game_state.current_lap > game_state.total_laps

    return {
        "started": game_state.started,
        "player_count": len(game_state.players),
        "players": ranking,
        "segment": current_segment(segments),
        "segment_index": game_state.segment_index,
        "segment_count": len(segments),
        "active_player": asdict(get_active_player()) if get_active_player() else None,
        "active_has_target": bool(game_state.active_position_index > 0),
        "last_duel": game_state.last_duel,
        "weather": game_state.weather,
        "weather_dex_penalty": weather_dex_penalty(),
        "lap": game_state.current_lap,
        "total_laps": game_state.total_laps,
        "finished": finished,
        "track": track_data(segments),
        "circuit_file": str(CIRCUIT_FILE),
        "final_stats": final_stats() if finished and game_state.started else None,
    }


@app.get("/")
def index() -> str:
    return render_template("index.html")


@app.get("/circuit-image")
def circuit_image() -> Any:
    if CIRCUIT_PNG.exists():
        return send_file(CIRCUIT_PNG)
    return ("circuit.png introuvable", 404)


@app.get("/api/state")
def get_state() -> Any:
    return jsonify(serialize_state())


@app.post("/api/restart")
def restart() -> Any:
    global game_state
    game_state = GameState()


PLAYER_COLORS = [
    "#ef4444", "#3b82f6", "#22c55e", "#f59e0b", "#a855f7",
    "#06b6d4", "#f97316", "#84cc16", "#e879f9", "#14b8a6",
]


def color_for_player(idx: int) -> str:
    return PLAYER_COLORS[(idx - 1) % len(PLAYER_COLORS)]
    return jsonify(serialize_state())


@app.post("/api/start")
def start_game() -> Any:
    payload = request.get_json(silent=True) or {}
    players_payload = payload.get("players", [])
    if not (3 <= len(players_payload) <= 10):
        return jsonify({"error": "Il faut entre 3 et 10 joueurs."}), 400

    total_laps = max(1, min(20, safe_int(payload.get("total_laps"), 5)))

    players: list[Player] = []
    for idx, row in enumerate(players_payload, start=1):
        name = (row.get("name") or "").strip()
        engine = safe_int(row.get("engine_power"), -1)
        downforce = safe_int(row.get("downforce"), -1)
        dex = safe_int(row.get("dexterity"), -1)
        aggr = safe_int(row.get("aggressiveness"), -1)
        if not name:
            return jsonify({"error": "Chaque joueur doit avoir un nom."}), 400
        if not (0 <= engine <= 20 and 0 <= downforce <= 20 and 0 <= dex <= 20 and 0 <= aggr <= 20):
            return jsonify({"error": "Toutes les stats doivent être entre 0 et 20."}), 400
        players.append(Player(player_id=idx, name=name, engine_power=engine, downforce=downforce, dexterity=dex, aggressiveness=aggr, color=color_for_player(idx)))

    game_state.players = players
    game_state.total_laps = total_laps
    game_state.current_lap = 1
    game_state.segment_index = 0
    game_state.active_position_index = 0
    game_state.last_duel = None
    game_state.weather = "dry"

    segments = load_circuit_segments()
    game_state.positions = qualification_order(segments) if segments else [p.player_id for p in players]

    return jsonify(serialize_state())


@app.post("/api/weather")
def set_weather() -> Any:
    payload = request.get_json(silent=True) or {}
    weather = payload.get("weather")
    if weather not in {"dry", "windy", "light_rain", "storm"}:
        return jsonify({"error": "Météo invalide."}), 400
    game_state.weather = weather
    return jsonify(serialize_state())


@app.post("/api/next-turn")
def next_turn() -> Any:
    state = serialize_state()
    if not state["started"]:
        return jsonify({"error": "La course n'est pas démarrée."}), 400
    if state["finished"]:
        return jsonify({"error": "Course terminée."}), 400
    active = get_active_player()
    if active is not None:
        game_state.last_duel = {"info": f"{active.name} passe son tour."}
    advance_turn()
    return jsonify(serialize_state())


@app.post("/api/action")
def action() -> Any:
    state = serialize_state()
    if not state["started"]:
        return jsonify({"error": "La course n'est pas démarrée."}), 400
    if state["finished"]:
        return jsonify({"error": "Course terminée."}), 400

    payload = request.get_json(silent=True) or {}
    action_type = payload.get("type")

    segments = load_circuit_segments()
    segment = current_segment(segments)
    if segment is None:
        return jsonify({"error": "Segment introuvable."}), 400

    attacker = get_active_player()
    if attacker is None:
        advance_turn()
        return jsonify(serialize_state())

    if action_type == "prepare":
        attacker.prep_defense_bonus = 2
        consume_turn(attacker, segment, action_wear=0.8)
        game_state.last_duel = {"info": f"{attacker.name} se prépare (+2 défense)."}
        advance_turn()
        return jsonify(serialize_state())

    if action_type == "next_section_bonus":
        attacker.next_section_bonus += 1
        consume_turn(attacker, segment, action_wear=0.8)
        game_state.last_duel = {"info": f"{attacker.name} prépare la prochaine section (+1 attaque)."}
        advance_turn()
        return jsonify(serialize_state())

    if action_type == "pass":
        consume_turn(attacker, segment, action_wear=0.6)
        game_state.last_duel = {"info": f"{attacker.name} passe cette section."}
        advance_turn()
        return jsonify(serialize_state())

    if action_type not in {"overtake", "dangerous_overtake"}:
        return jsonify({"error": "Action inconnue."}), 400

    if game_state.active_position_index == 0:
        consume_turn(attacker, segment, action_wear=0.6)
        game_state.last_duel = {"info": f"{attacker.name} est leader: pas de cible, tour passé."}
        advance_turn()
        return jsonify(serialize_state())

    defender = get_target_player()
    if defender is None:
        return jsonify({"error": "Défenseur indisponible."}), 400

    dangerous = action_type == "dangerous_overtake"

    atk_mod, atk_notes = attack_modifier(attacker, segment, dangerous)
    def_mod, def_notes = defense_modifier(defender)
    atk_roll = random.randint(1, 20)
    def_roll = random.randint(1, 20)

    risk_triggered, risk_note = dangerous_risk_penalty(attacker, dangerous)

    atk_total = atk_roll + atk_mod
    def_total = def_roll + def_mod
    success = atk_total > def_total

    if success:
        attacker.successful_overtakes += 1
        i = game_state.active_position_index
        game_state.positions[i - 1], game_state.positions[i] = game_state.positions[i], game_state.positions[i - 1]
        game_state.active_position_index -= 1
    else:
        defender.successful_defenses += 1

    defender.prep_defense_bonus = 0
    consume_turn(attacker, segment, action_wear=2.2 if dangerous else 1.6)

    game_state.last_duel = {
        "attacker": attacker.name,
        "defender": defender.name,
        "attack_roll": atk_roll,
        "attack_mod": atk_mod,
        "attack_total": atk_total,
        "attack_notes": atk_notes,
        "defense_roll": def_roll,
        "defense_mod": def_mod,
        "defense_total": def_total,
        "defense_notes": def_notes,
        "dangerous": dangerous,
        "risk_triggered": risk_triggered,
        "risk_note": risk_note,
        "success": success,
    }

    advance_turn()
    return jsonify(serialize_state())


def run_local_server() -> None:
    app.run(host="127.0.0.1", port=5000, debug=True, use_reloader=False)


if __name__ == "__main__":
    run_local_server()
